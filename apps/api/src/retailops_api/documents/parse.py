"""Read a supplier sheet from CSV or XLSX bytes.

Headers are normalised through the canonical alias map. Numeric cells are
parsed; blanks become ``None``. Each data row keeps the source row number
(1-based, matching the file: the header is row 1). Cell-level problems are
returned as parse issues so a reviewer can fix the file; they do not abort
the rest of the sheet.
"""

from __future__ import annotations

import csv
import io
from collections.abc import Sequence
from decimal import Decimal, InvalidOperation

from retailops_api.documents.schema import (
    DECIMAL_COLUMNS,
    INTEGER_COLUMNS,
    TEXT_COLUMNS,
    cell_values,
    map_headers,
    parse_decimal,
    parse_int,
)
from retailops_api.documents.types import ParseError, ParseIssue, ParseResult, SupplierSheetRow

CSV_MEDIA_TYPE = "text/csv"
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PDF_MEDIA_TYPE = "application/pdf"

MEDIA_TYPES_BY_SUFFIX = {
    ".csv": CSV_MEDIA_TYPE,
    ".xlsx": XLSX_MEDIA_TYPE,
    ".pdf": PDF_MEDIA_TYPE,
}


def detect_media_type(filename: str, media_type: str | None = None) -> str:
    if media_type:
        return media_type
    suffix = ""
    lowered = filename.rsplit("/", 1)[-1].rsplit("\\", 1)[-1].lower()
    if "." in lowered:
        suffix = "." + lowered.rsplit(".", 1)[-1]
    detected = MEDIA_TYPES_BY_SUFFIX.get(suffix)
    if detected is None:
        raise ParseError(
            f"cannot detect media type from filename {filename!r}",
            issues=[
                ParseIssue(
                    code="unsupported_media_type",
                    message=f"cannot detect media type from filename {filename!r}",
                )
            ],
        )
    return detected


def parse_supplier_sheet(
    data: bytes,
    *,
    media_type: str,
    filename: str = "document",
) -> ParseResult:
    """Dispatch on media type. Image-only PDFs are not handled here."""
    if not data:
        raise ParseError(
            "document is empty",
            issues=[ParseIssue(code="empty_document", message="document is empty")],
        )
    if media_type == CSV_MEDIA_TYPE:
        return parse_csv(data)
    if media_type == XLSX_MEDIA_TYPE:
        return parse_xlsx(data)
    if media_type == PDF_MEDIA_TYPE:
        from retailops_api.documents.pdf import parse_text_pdf

        return parse_text_pdf(data)
    raise ParseError(
        f"unsupported media type {media_type!r} for {filename}",
        issues=[
            ParseIssue(
                code="unsupported_media_type",
                message=f"unsupported media type {media_type!r}",
            )
        ],
    )


def parse_csv(data: bytes) -> ParseResult:
    text = _decode_text(data)
    reader = csv.reader(io.StringIO(text))
    table = [list(row) for row in reader]
    return _parse_row_table(table, source_row_offset=1)


def parse_xlsx(data: bytes) -> ParseResult:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        sheet = workbook.active
        if sheet is None:
            raise ParseError(
                "workbook has no active sheet",
                issues=[ParseIssue(code="empty_document", message="workbook has no active sheet")],
            )
        physical: list[tuple[int, list[object]]] = []
        for row in sheet.iter_rows(values_only=False):
            if not row:
                continue
            number = row[0].row or 0
            values = [cell.value for cell in row]
            physical.append((number, values))
    finally:
        workbook.close()
    if not physical:
        raise ParseError(
            "spreadsheet has no rows",
            issues=[ParseIssue(code="empty_document", message="spreadsheet has no rows")],
        )
    # Keep Excel row numbers so findings point at the cell the supplier sees.
    header_index = _first_nonempty_index([values for _, values in physical])
    header_cells = physical[header_index][1]
    mapping = map_headers([_as_header(cell) for cell in header_cells])
    _require_mapped_headers(mapping)
    issues: list[ParseIssue] = []
    rows: list[SupplierSheetRow] = []
    for number, cells in physical[header_index + 1 :]:
        if _row_blank(cells):
            continue
        row, row_issues = _build_row(mapping, cells, row_number=number)
        rows.append(row)
        issues.extend(row_issues)
    return ParseResult(rows=tuple(rows), issues=tuple(issues))


def parse_cell_table(
    table: Sequence[Sequence[object]],
    *,
    source_row_offset: int = 1,
) -> ParseResult:
    """Parse an already-reconstructed grid of cells into a supplier sheet."""
    return _parse_row_table(table, source_row_offset=source_row_offset)


def parse_delimited_table(text: str, *, source_row_offset: int = 1) -> ParseResult:
    """Parse already-extracted tabular text (CSV / TSV / pipe / semicolon)."""
    sample = text.lstrip("\ufeff")
    lines = sample.splitlines()
    if not lines:
        raise ParseError(
            "document has no lines",
            issues=[ParseIssue(code="empty_document", message="document has no lines")],
        )
    dialect_source = "\n".join(lines[:8])
    try:
        dialect = csv.Sniffer().sniff(dialect_source, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(sample), dialect)
    table = [list(row) for row in reader]
    return _parse_row_table(table, source_row_offset=source_row_offset)


def _parse_row_table(table: Sequence[Sequence[object]], *, source_row_offset: int) -> ParseResult:
    if not table:
        raise ParseError(
            "document has no rows",
            issues=[ParseIssue(code="empty_document", message="document has no rows")],
        )
    header_index = _first_nonempty_index(table)
    mapping = map_headers([_as_header(cell) for cell in table[header_index]])
    _require_mapped_headers(mapping)
    issues: list[ParseIssue] = []
    rows: list[SupplierSheetRow] = []
    for offset, cells in enumerate(table[header_index + 1 :], start=1):
        if _row_blank(cells):
            continue
        row_number = source_row_offset + header_index + offset
        row, row_issues = _build_row(mapping, cells, row_number=row_number)
        rows.append(row)
        issues.extend(row_issues)
    return ParseResult(rows=tuple(rows), issues=tuple(issues))


def _build_row(
    mapping: dict[int, str],
    cells: Sequence[object],
    *,
    row_number: int,
) -> tuple[SupplierSheetRow, list[ParseIssue]]:
    raw = cell_values(mapping, list(cells))
    invalid: set[str] = set()
    issues: list[ParseIssue] = []
    parsed: dict[str, object] = {}
    for name, value in raw.items():
        if value is None:
            parsed[name] = None
            continue
        if name in TEXT_COLUMNS:
            parsed[name] = value
            continue
        try:
            if name in DECIMAL_COLUMNS:
                parsed[name] = parse_decimal(value)
            elif name in INTEGER_COLUMNS:
                parsed[name] = parse_int(value)
            else:
                parsed[name] = value
        except (InvalidOperation, ValueError):
            invalid.add(name)
            parsed[name] = None
            issues.append(
                ParseIssue(
                    code="malformed_value",
                    message=f"{name} {value!r} is not a usable number",
                    row_number=row_number,
                    field=name,
                )
            )
    row = SupplierSheetRow(
        row_number=row_number,
        supplier_sku=_as_optional_str(parsed.get("supplier_sku")),
        ean=_as_optional_str(parsed.get("ean")),
        description=_as_optional_str(parsed.get("description")),
        category=_as_optional_str(parsed.get("category")),
        cost=_as_optional_decimal(parsed.get("cost")),
        vat=_as_optional_decimal(parsed.get("vat")),
        case_pack=_as_optional_int(parsed.get("case_pack")),
        minimum_order_quantity=_as_optional_int(parsed.get("minimum_order_quantity")),
        lead_time_days=_as_optional_int(parsed.get("lead_time_days")),
        invalid_fields=frozenset(invalid),
    )
    return row, issues


def _require_mapped_headers(mapping: dict[int, str]) -> None:
    if mapping:
        return
    raise ParseError(
        "no recognised supplier-sheet columns in the header row",
        issues=[
            ParseIssue(
                code="unrecognised_header",
                message="no recognised supplier-sheet columns in the header row",
                row_number=1,
            )
        ],
    )


def _first_nonempty_index(table: Sequence[Sequence[object]]) -> int:
    for index, cells in enumerate(table):
        if not _row_blank(cells):
            return index
    raise ParseError(
        "document has no non-empty rows",
        issues=[ParseIssue(code="empty_document", message="document has no non-empty rows")],
    )


def _row_blank(cells: Sequence[object]) -> bool:
    return all(cell is None or str(cell).strip() == "" for cell in cells)


def _as_header(cell: object) -> str:
    return "" if cell is None else str(cell)


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def _as_optional_str(value: object) -> str | None:
    return value if isinstance(value, str) else None


def _as_optional_decimal(value: object) -> Decimal | None:
    return value if isinstance(value, Decimal) else None


def _as_optional_int(value: object) -> int | None:
    return value if isinstance(value, int) and not isinstance(value, bool) else None
